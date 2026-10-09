#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
法務リーガルチェック結果を Word 赤入れ風 .docx で生成するスクリプト。

使用例：
    python generate_review_docx.py \
        --input review_data.json \
        --output contract_review_result.docx

review_data.json のスキーマ例：
{
  "contract_title": "業務委託契約書（株式会社XX）",
  "category": "gyomu-itaku",
  "category_label": "業務委託契約書",
  "check_date": "2026-04-23",
  "reviewer": "AI予備チェック（legal-contract-review スキル）",
  "summary_top3": [
    "違約金条項が一方的に当社に不利",
    "契約終了後の競業期間が5年（上限3年推奨）",
    "損害賠償の相互上限が設定されていない"
  ],
  "findings": [
    {
      "no": 1,
      "article": "第○条（違約金）",
      "severity": "必須",
      "topic": "違約金条項",
      "current_text": "乙は、甲に対し、本契約違反時に違約金として……",
      "expected_comment": "私的制裁を科される謂れはありません……",
      "proposed_text": "（削除）或いは「但し、以下のいずれかに該当する場合は……」",
      "reference": "過去事例：業務委託契約で同種の指摘あり"
    }
  ],
  "next_actions": [
    "先方に違約金の削除交渉",
    "契約終了後期間を3年以内に短縮",
    "損害賠償の相互上限（12ヶ月分）を追加"
  ]
}
"""

import argparse
import json
import sys
from pathlib import Path
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("python-docx is required. Install: pip install python-docx", file=sys.stderr)
    sys.exit(1)


# ============== カラー定義 ==============
COLOR_BLACK = RGBColor(0x00, 0x00, 0x00)
COLOR_RED = RGBColor(0xB8, 0x20, 0x2E)  # ライトアップ企業カラー
COLOR_NAVY = RGBColor(0x1F, 0x3A, 0x5F)
COLOR_GRAY = RGBColor(0x60, 0x60, 0x60)
COLOR_LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)

# 指摘区分の色
SEVERITY_COLOR = {
    "必須": RGBColor(0xC0, 0x00, 0x00),      # 濃い赤
    "推奨": RGBColor(0xE0, 0x7B, 0x00),      # オレンジ
    "確認": RGBColor(0x00, 0x66, 0xCC),      # 青
}


def set_cell_shading(cell, color_hex):
    """セルに背景色を設定"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def set_cell_borders(cell, color_hex="808080", size="4"):
    """セルに枠線を設定"""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color_hex)
        tc_borders.append(border)
    tc_pr.append(tc_borders)


def add_run(paragraph, text, *, bold=False, color=None, size_pt=None, font_name=None):
    """段落に装飾つきの run を追加"""
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    if color is not None:
        run.font.color.rgb = color
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if font_name is None:
        font_name = "游ゴシック"
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    return run


def add_heading(doc, text, *, level=1, color=None):
    """見出しを追加（色付き可）"""
    p = doc.add_paragraph()
    if level == 0:
        size = 22
    elif level == 1:
        size = 16
    elif level == 2:
        size = 13
    else:
        size = 11
    run = add_run(p, text, bold=True, color=color or COLOR_NAVY, size_pt=size)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_separator(doc):
    """水平罫線的な段落を追加"""
    p = doc.add_paragraph()
    add_run(p, "─" * 48, color=COLOR_GRAY, size_pt=9)
    p.paragraph_format.space_after = Pt(6)
    return p


def build_cover(doc, data):
    """表紙セクション"""
    # タイトル
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "リーガルチェック結果（AI予備チェック）",
            bold=True, color=COLOR_RED, size_pt=22)
    p.paragraph_format.space_after = Pt(24)

    # 契約書名
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, data.get("contract_title", "（契約書名未設定）"),
            bold=True, color=COLOR_BLACK, size_pt=16)
    p.paragraph_format.space_after = Pt(18)

    # メタ情報テーブル
    tbl = doc.add_table(rows=4, cols=2)
    tbl.autofit = False
    widths = [Cm(4.5), Cm(12.5)]

    meta = [
        ("契約書種別", data.get("category_label", data.get("category", "未判定"))),
        ("チェック日", data.get("check_date", datetime.now().strftime("%Y-%m-%d"))),
        ("実施者", data.get("reviewer", "AI予備チェック")),
        ("最終確認", "法務（legal@writeup.co.jp）による確認が必須"),
    ]
    for i, (label, value) in enumerate(meta):
        row = tbl.rows[i]
        for j, cell in enumerate(row.cells):
            cell.width = widths[j]
            set_cell_borders(cell, "B8B8B8")
            cell.paragraphs[0].paragraph_format.space_before = Pt(2)
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        set_cell_shading(row.cells[0], "F4E8E8")
        add_run(row.cells[0].paragraphs[0], label, bold=True, color=COLOR_NAVY, size_pt=10)
        add_run(row.cells[1].paragraphs[0], value, size_pt=10)

    doc.add_paragraph()

    # 免責ボックス
    p = doc.add_paragraph()
    add_run(p, "■ 免責事項", bold=True, color=COLOR_RED, size_pt=11)

    disclaimer = (
        "本ドキュメントは AI（legal-contract-review スキル）による"
        "予備チェック結果です。最終的な契約締結判断は必ず法務担当"
        "（legal@writeup.co.jp）の確認を経てください。"
        "AI の提案は過去のパターンに基づく機械的な抽出であり、"
        "個別具体の交渉状況・事業リスクを完全に反映したものでは"
        "ありません。"
    )
    p = doc.add_paragraph()
    add_run(p, disclaimer, color=COLOR_GRAY, size_pt=10)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(18)


def build_summary_top3(doc, data):
    """トップ3論点"""
    add_heading(doc, "■ 特に注意すべきトップ3論点", level=2, color=COLOR_RED)
    top3 = data.get("summary_top3", [])
    if not top3:
        p = doc.add_paragraph()
        add_run(p, "（該当なし。指摘論点サマリー表を参照してください）",
                color=COLOR_GRAY, size_pt=10)
        return

    for i, item in enumerate(top3, 1):
        p = doc.add_paragraph()
        add_run(p, f"{i}. ", bold=True, color=COLOR_RED, size_pt=11)
        add_run(p, item, size_pt=11)
        p.paragraph_format.space_after = Pt(3)
    doc.add_paragraph()


def build_summary_table(doc, data):
    """指摘サマリー表"""
    add_heading(doc, "■ 指摘論点サマリー", level=2, color=COLOR_NAVY)

    findings = data.get("findings", [])
    if not findings:
        p = doc.add_paragraph()
        add_run(p, "（指摘事項なし）", color=COLOR_GRAY, size_pt=10)
        return

    tbl = doc.add_table(rows=1 + len(findings), cols=5)
    tbl.autofit = False
    widths = [Cm(1.2), Cm(3.5), Cm(2.0), Cm(5.5), Cm(5.0)]
    headers = ["No.", "該当条項", "区分", "指摘論点", "改定方針（概要）"]

    # ヘッダ
    header_row = tbl.rows[0]
    for i, cell in enumerate(header_row.cells):
        cell.width = widths[i]
        set_cell_shading(cell, "1F3A5F")
        set_cell_borders(cell, "808080")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, headers[i], bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size_pt=10)

    # 本体
    for row_idx, f in enumerate(findings, 1):
        row = tbl.rows[row_idx]
        severity = f.get("severity", "確認")
        sev_color = SEVERITY_COLOR.get(severity, COLOR_GRAY)

        values = [
            str(f.get("no", row_idx)),
            f.get("article", "-"),
            severity,
            f.get("topic", "-"),
            (f.get("proposed_text", "")[:80] + "...") if len(f.get("proposed_text", "")) > 80 else f.get("proposed_text", "-"),
        ]
        for i, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = widths[i]
            set_cell_borders(cell, "B8B8B8")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_run(p, value, bold=True, size_pt=10)
            elif i == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_run(p, value, bold=True, color=sev_color, size_pt=10)
            else:
                add_run(p, value, size_pt=10)

    doc.add_paragraph()


def build_detail_section(doc, data):
    """詳細レビュー"""
    add_heading(doc, "■ 詳細レビュー", level=1, color=COLOR_RED)

    findings = data.get("findings", [])
    if not findings:
        return

    for f in findings:
        # 見出し行
        p = doc.add_paragraph()
        severity = f.get("severity", "確認")
        sev_color = SEVERITY_COLOR.get(severity, COLOR_GRAY)

        add_run(p, f"【{severity}】", bold=True, color=sev_color, size_pt=12)
        add_run(p, " ", size_pt=12)
        add_run(p, f"No.{f.get('no', '?')}  {f.get('article', '')}",
                bold=True, color=COLOR_NAVY, size_pt=12)
        add_run(p, "  ", size_pt=12)
        add_run(p, f.get("topic", ""), bold=True, color=COLOR_BLACK, size_pt=12)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)

        # 現行条文
        p = doc.add_paragraph()
        add_run(p, "  ▸ 現行条文（引用）：", bold=True, color=COLOR_GRAY, size_pt=10)

        current = f.get("current_text", "").strip()
        if current:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            add_run(p, current, color=COLOR_BLACK, size_pt=10)
            p.paragraph_format.space_after = Pt(4)

        # 想定される指摘
        p = doc.add_paragraph()
        add_run(p, "  ▸ 想定される法務指摘：", bold=True, color=COLOR_RED, size_pt=10)

        comment = f.get("expected_comment", "").strip()
        if comment:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            add_run(p, comment, color=COLOR_BLACK, size_pt=10)
            p.paragraph_format.space_after = Pt(4)

        # 理想的な改定文
        p = doc.add_paragraph()
        add_run(p, "  ▸ 理想的な改定文：", bold=True, color=COLOR_RED, size_pt=10)

        proposed = f.get("proposed_text", "").strip()
        if proposed:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            add_run(p, proposed, color=RGBColor(0xC0, 0x00, 0x00), size_pt=10)
            p.paragraph_format.space_after = Pt(4)

        # 参考
        reference = f.get("reference", "").strip()
        if reference:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            add_run(p, f"  ※ 参考：{reference}", color=COLOR_GRAY, size_pt=9)
            p.paragraph_format.space_after = Pt(4)

        add_separator(doc)


def build_next_actions(doc, data):
    """次のアクション"""
    add_heading(doc, "■ 次のアクション", level=2, color=COLOR_NAVY)

    actions = data.get("next_actions", [])
    if not actions:
        p = doc.add_paragraph()
        add_run(p, "（特記事項なし。指摘論点に従って交渉を進めてください）",
                color=COLOR_GRAY, size_pt=10)
        return

    for i, a in enumerate(actions, 1):
        p = doc.add_paragraph()
        add_run(p, f"{i}. ", bold=True, color=COLOR_NAVY, size_pt=11)
        add_run(p, a, size_pt=11)


def build_footer(doc):
    """フッター免責"""
    doc.add_paragraph()
    add_separator(doc)
    p = doc.add_paragraph()
    add_run(p, "※ ", color=COLOR_RED, size_pt=9)
    add_run(
        p,
        "本ドキュメントは AI による予備チェック結果です。"
        "最終的な契約締結判断は必ず法務担当（legal@writeup.co.jp）の確認を経てください。",
        color=COLOR_GRAY, size_pt=9,
    )


def build_document(data, output_path):
    """メインのドキュメント組み立て"""
    doc = Document()

    # ページ設定
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    # デフォルトフォント設定
    style = doc.styles["Normal"]
    style.font.name = "游ゴシック"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "游ゴシック")

    # セクション構築
    build_cover(doc, data)
    doc.add_page_break()
    build_summary_top3(doc, data)
    build_summary_table(doc, data)
    doc.add_page_break()
    build_detail_section(doc, data)
    build_next_actions(doc, data)
    build_footer(doc)

    doc.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="法務リーガルチェック結果をWord文書化")
    parser.add_argument("--input", "-i", required=True, help="review_data.json のパス")
    parser.add_argument("--output", "-o", required=True, help="出力 .docx のパス")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"ERROR: input file not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    with input_path.open("r", encoding="utf-8") as fp:
        data = json.load(fp)

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    build_document(data, str(output_path))
    print(f"OK: {output_path}")


if __name__ == "__main__":
    main()
